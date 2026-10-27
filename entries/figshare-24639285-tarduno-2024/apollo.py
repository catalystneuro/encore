"""Loading and measures for the Apollo magnetometer data of Figshare item 24639285, version 1.

Tarduno, Cottrell & Zhou (2024), Scientific Data, doi:10.1038/s41597-024-03626-0. Dataset
doi:10.6084/m9.figshare.24639285.v1 (the only version). Files are streamed from Figshare's
download service into memory; nothing is written to disk.

The deposit holds one workbook of 22 sheets for the Apollo 64455 glass subsamples
(64455lunar_cryodata.xlsx) and one workbook per single crystal (12 files). Each sheet has columns
sample, DEMAG (step), moment (emu), declination and inclination (degrees), one row per
measurement, with repeated rows for repeated measurements of a step. Steps: 0 is the NRM, "<x>mT"
an AF step, "<T>" or "<T>.0" field-off, "<T>.1" field-on, "<T>.2" a pTRM check (glass), "<T>.3"
the zero-field alteration check and "<T>.4" the 40 uT field-on step (crystals). The glass sheets
give the laboratory field in a footnote line ("Laboratory Field: 18 microTesla").

Where the parameters come from:

- TRM gained at a temperature is the vector difference of the field-on and field-off
  measurements (Methods, Thermal measurements, paragraph 3: "The TRM is the calculated vector
  subtraction of field-on and field-off steps"); NRM is the moment of the field-off measurement
  (claims file, opening comment).
- Repeated measurements of one step are combined as the mean vector (`how="mean"`) or the first
  measurement is used (`how="first"`). The paper does not say which Table 2 used; for the crystals
  it says the first measurements were used in ref. 1 (Background & Summary, paragraph 7). Both are
  run wherever they could matter (author, step 1.6).
- MAD is Kirschvink's (1980) principal component measure: free line (centred), line anchored to
  the origin (uncentred), or free line on the three-point running mean of the field-off vectors
  (claims file, opening comment, readings a to c). The running mean is centred, with the two end
  points averaged over the two points available, a standard choice.
- Great-circle MAD (fig2-changing-field, reading b) is Kirschvink's plane MAD with the plane
  through the origin, the standard form of a great circle; a standard choice.
- pTRM check percentage: the claims file's readings (a) |check pTRM - original pTRM| over the
  original pTRM, (b) the same difference over the length of the fitted Arai segment (DRAT-like,
  Selkin & Tauxe 2000), with the check's pTRM taken as the .2 vector minus (i) the field-off
  vector measured just before it or (ii) the field-off vector at the check's temperature. The
  difference is of magnitudes, the TRM-gained coordinate of the check on the Arai plot, where
  Figures 1 and 2 plot the checks (legends: "red triangles represent partial thermal remanent
  magnetization checks").
- Watson's (1956) test of randomness: p is the fraction of 200000 sets of n uniform random unit
  vectors whose resultant length is at least R (an exact Monte Carlo form of the test, with a
  fixed seed per call); the claims file names the test.
- Magnetometer threshold 8e-11 emu and the max intensity formula: Methods, Magnetometer
  measurements, and Technical Validation.
"""

from __future__ import annotations

import io
import json
import re
import urllib.request

import fsspec
import numpy as np
import openpyxl
import pandas as pd

FIGSHARE_ID = "24639285"
FIGSHARE_VERSION = 1
DATASET_DOI = "10.6084/m9.figshare.24639285.v1"
API_URL = f"https://api.figshare.com/v2/articles/{FIGSHARE_ID}/versions/{FIGSHARE_VERSION}"
GLASS_FILE = "64455lunar_cryodata.xlsx"
THRESHOLD_EMU = 8e-11


# --- reading -------------------------------------------------------------------

def file_list() -> list[dict]:
    """The pinned version's files (id, name, size, computed_md5) from the Figshare API."""
    with urllib.request.urlopen(API_URL) as r:
        meta = json.load(r)
    assert meta["version"] == FIGSHARE_VERSION and meta["doi"] == DATASET_DOI, (meta["version"], meta["doi"])
    return meta["files"]


def read_bytes(file_id: int) -> bytes:
    url = f"https://ndownloader.figshare.com/files/{file_id}"
    with fsspec.open(url, block_size=2**20) as f:
        return f.read()


def parse_sheet(ws) -> dict:
    """One sheet: its header, the measurement rows (in sheet order) and its footnote lines."""
    rows = list(ws.iter_rows(values_only=True))
    header = [h for h in rows[0] if h is not None]
    i_m = next(i for i, h in enumerate(rows[0]) if h is not None and str(h).startswith("Moment"))
    records, blank, notes = [], [], []
    for n, r in enumerate(rows[1:], start=2):
        if r[0] is not None and r[1] is not None:
            m, dec, inc = r[i_m], r[i_m + 1], r[i_m + 2]
            if m is None or dec is None or inc is None:
                blank.append(n)
                continue
            records.append({"row": n, "sample": str(r[0]), "step": str(r[1]),
                            "moment": float(m), "dec": float(dec), "inc": float(inc)})
        elif r[0] is not None:
            notes.append(str(r[0]).strip())
    df = pd.DataFrame(records)
    field = None
    for note in notes:
        hit = re.search(r"Laboratory Field:\s*([\d.]+)\s*micro", note)
        if hit:
            field = float(hit.group(1))
    return {"header": header, "data": df, "notes": notes, "lab_field_uT": field, "incomplete_rows": blank}


def load_all() -> tuple[list[dict], dict, dict]:
    """(file list, glass sheets by sheet name, crystal sheets by crystal name)."""
    files = file_list()
    glass, crystals = {}, {}
    for f in files:
        if not f["name"].endswith(".xlsx"):
            continue
        wb = openpyxl.load_workbook(io.BytesIO(read_bytes(f["id"])), data_only=True)
        f["sheets"] = wb.sheetnames
        for ws in wb.worksheets:
            sheet = parse_sheet(ws)
            sheet["file"] = f["name"]
            if f["name"] == GLASS_FILE:
                glass[ws.title] = sheet
            else:
                crystals[f["name"].split(".")[0]] = sheet
    return files, glass, crystals


# --- steps and vectors ---------------------------------------------------------

def step_kind(step: str) -> tuple[str, float | None]:
    """('nrm' | 'af' | 'off' | 'on' | 'check' | 'alt' | 'on40', temperature or AF field)."""
    if step.endswith("mT"):
        return "af", float(step[:-2])
    value = float(step)
    if value == 0:
        return "nrm", 0.0
    temp = float(int(value))
    code = round((value - temp) * 10)
    return {0: "off", 1: "on", 2: "check", 3: "alt", 4: "on40"}[code], temp


def to_xyz(moment, dec, inc) -> np.ndarray:
    d, i = np.radians(np.asarray(dec, float)), np.radians(np.asarray(inc, float))
    m = np.asarray(moment, float)
    return np.stack([m * np.cos(i) * np.cos(d), m * np.cos(i) * np.sin(d), m * np.sin(i)], axis=-1)


def to_dir(v) -> tuple[float, float, float]:
    v = np.asarray(v, float)
    m = np.linalg.norm(v)
    return float(np.degrees(np.arctan2(v[1], v[0])) % 360), float(np.degrees(np.arcsin(v[2] / m))), float(m)


def angle(a, b) -> float:
    a, b = np.asarray(a, float), np.asarray(b, float)
    c = np.dot(a, b) / np.linalg.norm(a) / np.linalg.norm(b)
    return float(np.degrees(np.arccos(np.clip(c, -1, 1))))


def step_table(df: pd.DataFrame, how: str = "mean") -> pd.DataFrame:
    """One row per step label in order of first appearance, with its vector (mean or first)."""
    out = []
    for step, g in df.groupby("step", sort=False):
        xyz = to_xyz(g.moment, g.dec, g.inc)
        v = xyz.mean(axis=0) if how == "mean" else xyz[0]
        kind, t = step_kind(step)
        out.append({"step": step, "kind": kind, "t": t, "n": len(g), "x": v[0], "y": v[1], "z": v[2],
                    "moment": float(np.linalg.norm(v)), "order": int(g.row.min())})
    return pd.DataFrame(out).sort_values("order").reset_index(drop=True)


def vec(row) -> np.ndarray:
    return np.array([row["x"], row["y"], row["z"]])


def scaled_copies(values, lo: float = 0.90, hi: float = 1.10, tol: float = 1e-6) -> bool:
    """Whether every value equals one base value times a factor k/100 in [lo, hi] (to tol, relative)."""
    v = np.asarray(values, float)
    if len(v) < 3 or len(np.unique(v)) < 2 or np.any(v == 0):
        return False
    for v0 in v:
        for f0 in np.arange(round(lo * 100), round(hi * 100) + 1) / 100:
            r = v / (v0 / f0) * 100
            if np.all(np.abs(r - np.round(r)) < tol * np.abs(r)) and np.all((r >= lo * 100 - 1e-6) & (r <= hi * 100 + 1e-6)):
                return True
    return False


def synthetic_steps(df: pd.DataFrame) -> list[str]:
    """Steps whose repeated rows are all one row scaled by factors of 0.90 to 1.10, in moment,
    declination and inclination alike."""
    return [step for step, g in df.groupby("step", sort=False)
            if all(scaled_copies(g[c].abs()) for c in ("moment", "dec", "inc"))]


# --- Arai and directions -------------------------------------------------------

def arai(df: pd.DataFrame, how: str = "mean") -> pd.DataFrame:
    """Paired thermal steps: NRM (moment of the field-off vector) and TRM gained (|on - off|)."""
    st = step_table(df, how)
    offs = st[st.kind == "off"].set_index("t")
    ons = st[st.kind == "on"].set_index("t")
    rows = []
    for t in sorted(set(offs.index) & set(ons.index)):
        off, on = vec(offs.loc[t]), vec(ons.loc[t])
        rows.append({"t": t, "nrm": float(np.linalg.norm(off)), "trm": float(np.linalg.norm(on - off)),
                     "x": off[0], "y": off[1], "z": off[2]})
    return pd.DataFrame(rows)


def in_range(a: pd.DataFrame, t0: float, t1: float) -> pd.DataFrame:
    return a[(a.t >= t0) & (a.t <= t1)].reset_index(drop=True)


def r2(x, y) -> float:
    return float(np.corrcoef(x, y)[0, 1] ** 2)


def slope(x, y) -> float:
    return float(np.polyfit(x, y, 1)[0])


def mad_line(points, anchored: bool = False) -> float:
    X = np.asarray(points, float)
    if not anchored:
        X = X - X.mean(axis=0)
    w = np.sort(np.linalg.eigvalsh(X.T @ X))[::-1]
    w = np.clip(w, 0, None)
    return float(np.degrees(np.arctan(np.sqrt((w[1] + w[2]) / w[0]))))


def mad_plane(points) -> float:
    """Great-circle MAD, plane through the origin."""
    X = np.asarray(points, float)
    w = np.clip(np.sort(np.linalg.eigvalsh(X.T @ X))[::-1], 0, None)
    return float(np.degrees(np.arctan(np.sqrt(w[2] / w[1] + w[2] / w[0]))))


def demag_sequence(df: pd.DataFrame, how: str = "mean") -> pd.DataFrame:
    """The field-off vectors in demagnetization order: the last NRM or AF step, then the thermal
    field-off steps by temperature."""
    st = step_table(df, how)
    start = st[st.kind.isin(["nrm", "af"])].iloc[[-1]]
    offs = st[st.kind == "off"].sort_values("t")
    return pd.concat([start, offs]).reset_index(drop=True)


def running_mean3(X) -> np.ndarray:
    X = np.asarray(X, float)
    out = np.empty_like(X)
    for i in range(len(X)):
        out[i] = X[max(0, i - 1): i + 2].mean(axis=0)
    return out


def mads(df: pd.DataFrame, t0: float, t1: float, how: str = "mean") -> dict:
    """MAD of the field-off vectors from t0 to t1 under readings (a) free, (b) anchored, (c) free
    on the three-point running mean."""
    seq = demag_sequence(df, how)
    sel = ((seq.kind == "off") & (seq.t >= t0) & (seq.t <= t1)).to_numpy()
    X = seq[["x", "y", "z"]].to_numpy()
    return {"a_free": mad_line(X[sel]), "b_anchored": mad_line(X[sel], anchored=True),
            "c_smoothed": mad_line(running_mean3(X)[sel])}


def ptrm_checks(df: pd.DataFrame, how: str = "mean", seg: tuple[float, float] | None = None) -> pd.DataFrame:
    """Each pTRM check, as a percentage under readings (a)/(b) x (i)/(ii)."""
    st = step_table(df, how)
    offs = st[st.kind == "off"].set_index("t")
    ons = st[st.kind == "on"].set_index("t")
    a = arai(df, how)
    length = np.nan
    if seg is not None:
        s = in_range(a, *seg)
        if len(s) >= 2:
            b = slope(s.trm, s.nrm)
            dx = s.trm.max() - s.trm.min()
            length = float(np.hypot(dx, b * dx))
    rows = []
    for i, r in st[st.kind == "check"].iterrows():
        t = r.t
        prev = st.iloc[:i][st.iloc[:i].kind == "off"]
        ref_prev = vec(prev.iloc[-1])
        orig = np.linalg.norm(vec(ons.loc[t]) - vec(offs.loc[t]))
        for ref_name, ref in (("i_prev", ref_prev), ("ii_same", vec(offs.loc[t]))):
            chk = np.linalg.norm(vec(r) - ref)
            rows.append({"t": t, "ref": ref_name, "prev_off_t": float(prev.iloc[-1].t),
                         "check_trm": chk, "orig_trm": orig,
                         "a_pct": 100 * abs(chk - orig) / orig,
                         "b_pct": 100 * abs(chk - orig) / length if length == length else np.nan})
    return pd.DataFrame(rows, columns=["t", "ref", "prev_off_t", "check_trm", "orig_trm", "a_pct", "b_pct"])


# --- statistics on directions --------------------------------------------------

def fisher(xyz) -> dict:
    u = np.asarray(xyz, float)
    u = u / np.linalg.norm(u, axis=1, keepdims=True)
    n = len(u)
    R = float(np.linalg.norm(u.sum(axis=0)))
    mean = to_dir(u.sum(axis=0))
    k = (n - 1) / (n - R) if n > 1 and n - R > 0 else np.nan
    a95 = np.nan
    if n > 1:
        c = 1 - (n - R) / R * ((1 / 0.05) ** (1 / (n - 1)) - 1)
        a95 = float(np.degrees(np.arccos(c))) if -1 <= c <= 1 else 180.0
    return {"n": n, "R": R, "dec": mean[0], "inc": mean[1], "k": k, "a95": a95}


def watson_p(n: int, R: float, seed: int, nsim: int = 200_000) -> float:
    """P(resultant of n uniform unit vectors >= R)."""
    if n < 2:
        return np.nan
    rng = np.random.default_rng(seed)
    v = rng.normal(size=(nsim, n, 3))
    v /= np.linalg.norm(v, axis=2, keepdims=True)
    return float((np.linalg.norm(v.sum(axis=1), axis=1) >= R - 1e-12).mean())


def max_intensity(m590_emu: float, lab_field_uT: float = 20.0) -> float:
    """Technical Validation: threshold / TRM at 590 C x lab field."""
    return THRESHOLD_EMU / m590_emu * lab_field_uT
